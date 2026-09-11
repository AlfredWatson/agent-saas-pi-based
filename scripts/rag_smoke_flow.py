"""End-to-end HTTP acceptance for the multi-tenant RAG backend.

Run the Gateway, RAG worker, Redis, PostgreSQL and ``rag_mock_model.py`` first.
The script creates isolated smoke data and deletes both knowledge bases at the end.
"""

from __future__ import annotations

import asyncio
import os
import sys
import time
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import httpx
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation
from sqlalchemy import delete, func, select

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "gateway"))

from app.db.session import SessionLocal, engine  # noqa: E402
from app.db.models import User, Workspace  # noqa: E402
from app.rag.cache import RagCache  # noqa: E402
from app.rag.models import (  # noqa: E402
    Chunk,
    ChunkVector,
    GraphArtifact,
    GraphEdge,
    GraphEvidence,
    GraphNode,
    ProcessingJob,
    RagDocument,
)

GATEWAY = os.getenv("RAG_SMOKE_GATEWAY", "http://127.0.0.1:21996/api/v1")
MODEL = os.getenv("RAG_SMOKE_MODEL", "http://127.0.0.1:22001")


def pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, item in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode() + item + b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(result)


def sample_files() -> list[tuple[str, tuple[str, bytes, str]]]:
    repeated = "Acme uses RAG for reliable retrieval. " * 120
    word_buffer = BytesIO()
    word = WordDocument()
    word.add_heading("Acme Word", 1)
    word.add_paragraph(repeated)
    word.save(word_buffer)

    excel_buffer = BytesIO()
    workbook = Workbook()
    workbook.active.title = "Acme Data"
    for index in range(80):
        workbook.active.append([index, "Acme", "RAG retrieval"])
    workbook.save(excel_buffer)

    ppt_buffer = BytesIO()
    presentation = Presentation()
    for index in range(8):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Acme RAG slide {index} " + repeated[:240]
    presentation.save(ppt_buffer)

    return [
        ("files", ("sample.pdf", pdf_bytes(repeated[:500]), "application/pdf")),
        (
            "files",
            (
                "sample.docx",
                word_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        ),
        (
            "files",
            (
                "sample.xlsx",
                excel_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        ),
        (
            "files",
            (
                "sample.pptx",
                ppt_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            ),
        ),
        ("files", ("sample.md", ("# Acme\n" + repeated).encode(), "text/markdown")),
    ]


def require(response: httpx.Response) -> dict:
    response.raise_for_status()
    return response.json() if response.content else {}


def reset_model(client: httpx.Client, **body) -> dict:
    return require(client.post(f"{MODEL}/control/reset", json=body))


def poll_documents(
    client: httpx.Client,
    headers: dict[str, str],
    root: str,
    document_ids: list[str],
    stage: str,
    expected: str = "succeeded",
    timeout: float = 90,
) -> dict[str, dict]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        rows = require(client.get(f"{root}/documents", headers=headers))["items"]
        selected = {row["id"]: row for row in rows if row["id"] in document_ids}
        statuses = {row["stages"][stage]["status"] for row in selected.values()}
        if len(selected) == len(document_ids) and statuses == {expected}:
            return selected
        if expected == "succeeded" and "failed" in statuses:
            errors = {
                row["filename"]: row["stages"][stage]["error"]
                for row in selected.values()
                if row["stages"][stage]["status"] == "failed"
            }
            raise AssertionError(f"{stage} failed: {errors}")
        time.sleep(0.2)
    raise TimeoutError(f"timed out waiting for {stage}={expected}")


def poll_operation(
    client: httpx.Client, headers: dict[str, str], operation_id: str
) -> dict:
    for _ in range(450):
        operation = require(
            client.get(f"{GATEWAY}/rag/operations/{operation_id}", headers=headers)
        )
        if operation["status"] == "succeeded":
            return operation
        if operation["status"] == "failed":
            raise AssertionError(operation)
        time.sleep(0.2)
    raise TimeoutError("operation timed out")


async def database_counts(knowledge_base_id: UUID) -> dict[str, int]:
    models = {
        "documents": RagDocument,
        "chunks": Chunk,
        "vectors": ChunkVector,
        "graphs": GraphArtifact,
        "nodes": GraphNode,
        "edges": GraphEdge,
        "evidence": GraphEvidence,
        "jobs": ProcessingJob,
    }
    async with SessionLocal() as db:
        counts = {}
        for name, model in models.items():
            column = (
                model.knowledge_base_id
                if hasattr(model, "knowledge_base_id")
                else GraphArtifact.knowledge_base_id
            )
            statement = select(func.count()).select_from(model)
            if hasattr(model, "knowledge_base_id"):
                statement = statement.where(column == knowledge_base_id)
            else:
                statement = statement.join(
                    GraphArtifact, GraphArtifact.id == model.artifact_id
                ).where(GraphArtifact.knowledge_base_id == knowledge_base_id)
            counts[name] = int(await db.scalar(statement) or 0)
        return counts


async def compare_database_counts(
    source_id: UUID, target_id: UUID
) -> tuple[dict[str, int], dict[str, int]]:
    try:
        return await database_counts(source_id), await database_counts(target_id)
    finally:
        await engine.dispose()


async def cache_key_count(kb_id: UUID, document_id: UUID) -> int:
    cache = RagCache()
    try:
        return int(await cache.client.scard(cache.index_key(kb_id, document_id)))
    finally:
        await cache.close()


async def cleanup_users(emails: list[str]) -> None:
    try:
        async with SessionLocal() as db:
            user_ids = list(
                (await db.scalars(select(User.id).where(User.email.in_(emails)))).all()
            )
            if user_ids:
                await db.execute(
                    delete(Workspace).where(Workspace.user_id.in_(user_ids))
                )
                await db.execute(delete(User).where(User.id.in_(user_ids)))
            await db.commit()
    finally:
        await engine.dispose()


def submit(
    client: httpx.Client,
    headers: dict[str, str],
    root: str,
    kind: str,
    document_ids: list[str],
) -> dict:
    return require(
        client.post(
            f"{root}/jobs/{kind}",
            headers=headers,
            json={"document_ids": document_ids},
        )
    )


def main() -> None:
    email = f"rag-smoke-{uuid4().hex[:12]}@example.com"
    password = "correct-horse-battery-staple"
    with httpx.Client(timeout=90, trust_env=False) as client:
        token = require(
            client.post(
                f"{GATEWAY}/auth/register",
                json={"email": email, "password": password},
            )
        )["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
        workspace_id = require(client.get(f"{GATEWAY}/workspaces", headers=headers))[
            "items"
        ][0]["id"]
        created = require(
            client.post(
                f"{GATEWAY}/workspaces/{workspace_id}/knowledge-bases",
                headers=headers,
                json={"name": f"rag-smoke-{uuid4().hex[:8]}"},
            )
        )
        kb_id = created["id"]
        root = f"{GATEWAY}/workspaces/{workspace_id}/knowledge-bases/{kb_id}"
        require(
            client.patch(
                root,
                headers=headers,
                json={
                    "chunking_strategy": "fixed",
                    "chunking_config": {
                        "max_token_size": 32,
                        "overlap_token_size": 4,
                        "split_by_character": "\n\n",
                    },
                    "chunking_concurrency": 2,
                    "embedding_concurrency": 1,
                    "graph_concurrency": 1,
                },
            )
        )
        documents = require(
            client.post(f"{root}/documents", headers=headers, files=sample_files())
        )["items"]
        assert len(documents) == 5
        document_ids = [item["id"] for item in documents]
        markdown_id = next(
            item["id"] for item in documents if item["filename"] == "sample.md"
        )

        model_body = {
            "protocol": "openai",
            "base_url": f"{MODEL}/v1",
            "api_key": "smoke-key",
            "model_name": "smoke-embedding",
        }
        require(client.put(f"{root}/embedding-model", headers=headers, json=model_body))
        model_body["model_name"] = "smoke-llm"
        require(client.put(f"{root}/llm-model", headers=headers, json=model_body))
        models = require(client.get(f"{root}/models", headers=headers))
        assert "smoke-key" not in str(models) and "api_key" not in str(models)

        submit(client, headers, root, "chunking", document_ids)
        poll_documents(client, headers, root, document_ids, "chunking")

        reset_model(client, embedding_fail_after=1)
        submit(client, headers, root, "vectorization", [markdown_id])
        poll_documents(
            client, headers, root, [markdown_id], "vectorization", expected="failed"
        )
        cached_after_failure = asyncio.run(
            cache_key_count(UUID(kb_id), UUID(markdown_id))
        )
        assert cached_after_failure >= 1
        reset_model(client)
        submit(client, headers, root, "vectorization", [markdown_id])
        poll_documents(client, headers, root, [markdown_id], "vectorization")
        submit(
            client,
            headers,
            root,
            "vectorization",
            [item for item in document_ids if item != markdown_id],
        )
        poll_documents(client, headers, root, document_ids, "vectorization")

        reset_model(client, chat_fail_after=1)
        submit(client, headers, root, "graph-extraction", [markdown_id])
        poll_documents(client, headers, root, [markdown_id], "graph", expected="failed")
        reset_model(client)
        submit(client, headers, root, "graph-extraction", [markdown_id])
        poll_documents(client, headers, root, [markdown_id], "graph")
        submit(
            client,
            headers,
            root,
            "graph-extraction",
            [item for item in document_ids if item != markdown_id],
        )
        poll_documents(client, headers, root, document_ids, "graph")

        for mode in ("vector", "hybrid"):
            result = require(
                client.post(
                    f"{root}/retrieve",
                    headers=headers,
                    json={"query": "Acme RAG", "mode": mode, "top_k": 5},
                )
            )
            assert result["items"]
        graph_result = require(
            client.post(
                f"{root}/retrieve",
                headers=headers,
                json={"query": "Acme", "mode": "graph", "top_k": 5},
            )
        )
        assert graph_result["nodes"] and graph_result["evidence"]

        source_graphs = require(client.get(f"{root}/graphs", headers=headers))["items"]
        document_graph_ids = [
            item["id"] for item in source_graphs if item["kind"] == "document"
        ]
        assert len(document_graph_ids) == 5
        merged = require(
            client.post(
                f"{root}/graphs:merge",
                headers=headers,
                json={
                    "name": "merged smoke graph",
                    "graph_ids": document_graph_ids[:2],
                },
            )
        )
        graph_ids_after_merge = {
            item["id"]
            for item in require(client.get(f"{root}/graphs", headers=headers))["items"]
        }
        assert merged["id"] in graph_ids_after_merge
        assert set(document_graph_ids) <= graph_ids_after_merge

        copied = require(
            client.post(
                f"{root}/copy",
                headers=headers,
                json={"name": f"rag-copy-{uuid4().hex[:8]}"},
            )
        )
        poll_operation(client, headers, copied["operation_id"])
        copy_id = copied["target_knowledge_base_id"]
        copy_root = f"{GATEWAY}/workspaces/{workspace_id}/knowledge-bases/{copy_id}"
        source_counts, copy_counts = asyncio.run(
            compare_database_counts(UUID(kb_id), UUID(copy_id))
        )
        assert source_counts == copy_counts, (source_counts, copy_counts)
        assert (
            len(require(client.get(f"{copy_root}/documents", headers=headers))["items"])
            == 5
        )
        assert (
            len(require(client.get(f"{copy_root}/graphs", headers=headers))["items"])
            == 6
        )

        locked_model = {
            "protocol": "openai",
            "base_url": f"{MODEL}/v1",
            "api_key": "smoke-key",
            "model_name": "another-embedding",
        }
        locked = client.put(
            f"{root}/embedding-model", headers=headers, json=locked_model
        )
        assert locked.status_code == 409

        copy_documents = require(client.get(f"{copy_root}/documents", headers=headers))[
            "items"
        ]
        cross_filter = client.post(
            f"{root}/retrieve",
            headers=headers,
            json={
                "query": "Acme",
                "mode": "vector",
                "document_ids": [copy_documents[0]["id"]],
            },
        )
        assert cross_filter.status_code == 422
        copy_graph = require(client.get(f"{copy_root}/graphs", headers=headers))[
            "items"
        ][0]
        cross_merge = client.post(
            f"{root}/graphs:merge",
            headers=headers,
            json={
                "name": "forbidden cross knowledge base merge",
                "graph_ids": [document_graph_ids[0], copy_graph["id"]],
            },
        )
        assert cross_merge.status_code == 422

        attacker_email = f"rag-attacker-{uuid4().hex[:12]}@example.com"
        attacker = require(
            client.post(
                f"{GATEWAY}/auth/register",
                json={
                    "email": attacker_email,
                    "password": password,
                },
            )
        )["access_token"]
        attacker_headers = {"Authorization": f"Bearer {attacker}"}
        assert client.get(root, headers=attacker_headers).status_code == 404
        assert (
            client.get(
                f"{root}/documents/{markdown_id}", headers=attacker_headers
            ).status_code
            == 404
        )
        assert (
            client.get(
                f"{GATEWAY}/rag/operations/{copied['operation_id']}",
                headers=attacker_headers,
            ).status_code
            == 404
        )

        require(
            client.delete(f"{root}/documents/{markdown_id}/vectors", headers=headers)
        )
        submit(client, headers, root, "vectorization", [markdown_id])
        poll_documents(client, headers, root, [markdown_id], "vectorization")

        deleted_document = next(item for item in documents if item["id"] != markdown_id)
        require(
            client.delete(f"{root}/documents/{deleted_document['id']}", headers=headers)
        )
        missing = client.get(
            f"{root}/documents/{deleted_document['id']}", headers=headers
        )
        assert missing.status_code == 404
        assert (
            asyncio.run(cache_key_count(UUID(kb_id), UUID(deleted_document["id"]))) == 0
        )

        for candidate in (kb_id, copy_id):
            operation = require(
                client.delete(
                    f"{GATEWAY}/workspaces/{workspace_id}/knowledge-bases/{candidate}",
                    headers=headers,
                )
            )
            poll_operation(client, headers, operation["operation_id"])

        asyncio.run(cleanup_users([email, attacker_email]))

        print(
            "PASS: five document types, model validation, three queues, cached resume, "
            "vector/hybrid/graph retrieval, graph merge, deep copy, re-vectorization, "
            "tenant isolation, and cascade deletion"
        )


if __name__ == "__main__":
    main()
