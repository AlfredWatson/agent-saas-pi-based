from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Protocol

from docx import Document as WordDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".ppt",
    ".pptx",
    ".md",
    ".markdown",
}


class UnsupportedDocumentError(ValueError):
    pass


class DocumentProcessor(Protocol):
    def parse(self, filename: str, content: bytes) -> list[Document]: ...


class DefaultDocumentProcessor:
    def parse(self, filename: str, content: bytes) -> list[Document]:
        extension = Path(filename).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentError(
                f"unsupported_document_type:{extension or 'none'}"
            )
        parsers = {
            ".pdf": self._pdf,
            ".doc": self._legacy_office,
            ".docx": self._word,
            ".xls": self._legacy_office,
            ".xlsx": self._excel,
            ".ppt": self._legacy_office,
            ".pptx": self._powerpoint,
            ".md": self._markdown,
            ".markdown": self._markdown,
        }
        documents = parsers[extension](filename, content)
        return [item for item in documents if item.page_content.strip()]

    @staticmethod
    def _legacy_office(filename: str, content: bytes) -> list[Document]:
        """Best-effort text extraction for legacy OLE Office files.

        The binary formats store user-visible strings as a mix of single-byte
        and UTF-16LE runs. Images and formatting records are intentionally
        ignored by the default backend.
        """
        candidates: list[tuple[int, str]] = []
        for match in re.finditer(rb"(?:[\x20-\x7e\x80-\xff]\x00){4,}", content):
            value = match.group().decode("utf-16le", errors="ignore").strip()
            if value:
                candidates.append((match.start(), value))
        unicode_run = re.compile(
            r"[A-Za-z0-9\u3400-\u9fff][A-Za-z0-9\u3400-\u9fff\s，。！？；：、,.!?;:'\"()（）_-]{3,}"
        )
        for offset in (0, 1):
            decoded = content[offset:].decode("utf-16le", errors="ignore")
            for match in unicode_run.finditer(decoded):
                candidates.append((offset + match.start() * 2, match.group().strip()))
        for match in re.finditer(rb"[\x20-\x7e]{4,}", content):
            value = match.group().decode("cp1252", errors="ignore").strip()
            if value:
                candidates.append((match.start(), value))
        candidates.sort(key=lambda item: item[0])
        seen: set[str] = set()
        text_runs: list[str] = []
        for _, value in candidates:
            normalized = " ".join(value.split())
            if normalized not in seen and any(
                character.isalnum() for character in normalized
            ):
                seen.add(normalized)
                text_runs.append(normalized)
        return [
            Document(
                page_content="\n".join(text_runs),
                metadata={
                    "source": filename,
                    "legacy_format": Path(filename).suffix.lower().lstrip("."),
                    "extraction": "best_effort_text",
                },
            )
        ]

    @staticmethod
    def _pdf(filename: str, content: bytes) -> list[Document]:
        reader = PdfReader(BytesIO(content))
        return [
            Document(
                page_content=page.extract_text() or "",
                metadata={"source": filename, "page": index + 1},
            )
            for index, page in enumerate(reader.pages)
        ]

    @staticmethod
    def _word(filename: str, content: bytes) -> list[Document]:
        word = WordDocument(BytesIO(content))
        blocks: list[Document] = []
        heading = ""
        buffer: list[str] = []

        def flush() -> None:
            text = "\n".join(buffer).strip()
            if text:
                blocks.append(
                    Document(
                        page_content=text,
                        metadata={"source": filename, "heading": heading},
                    )
                )
            buffer.clear()

        for paragraph in word.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                flush()
                heading = text
            else:
                buffer.append(text)
        for table in word.tables:
            rows = [
                "\t".join(cell.text.strip() for cell in row.cells) for row in table.rows
            ]
            if rows:
                buffer.append("\n".join(rows))
        flush()
        return blocks

    @staticmethod
    def _excel(filename: str, content: bytes) -> list[Document]:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        blocks: list[Document] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            part = 1
            for values in sheet.iter_rows(values_only=True):
                rendered = "\t".join(
                    "" if value is None else str(value) for value in values
                ).rstrip()
                if rendered:
                    rows.append(rendered)
                if len(rows) == 100:
                    blocks.append(
                        Document(
                            page_content="\n".join(rows),
                            metadata={
                                "source": filename,
                                "sheet": sheet.title,
                                "part": part,
                            },
                        )
                    )
                    rows = []
                    part += 1
            if rows:
                blocks.append(
                    Document(
                        page_content="\n".join(rows),
                        metadata={
                            "source": filename,
                            "sheet": sheet.title,
                            "part": part,
                        },
                    )
                )
        workbook.close()
        return blocks

    @staticmethod
    def _powerpoint(filename: str, content: bytes) -> list[Document]:
        presentation = Presentation(BytesIO(content))
        documents: list[Document] = []
        for index, slide in enumerate(presentation.slides):
            parts = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            documents.append(
                Document(
                    page_content="\n".join(parts),
                    metadata={"source": filename, "slide": index + 1},
                )
            )
        return documents

    @staticmethod
    def _markdown(filename: str, content: bytes) -> list[Document]:
        text = content.decode("utf-8-sig")
        headings: list[str] = []
        current: list[str] = []
        blocks: list[Document] = []

        def flush() -> None:
            body = "\n".join(current).strip()
            if body:
                blocks.append(
                    Document(
                        page_content=body,
                        metadata={"source": filename, "heading": " / ".join(headings)},
                    )
                )
            current.clear()

        for line in text.splitlines():
            match = re.match(r"^(#{1,6})\s+(.+)$", line)
            if match:
                flush()
                level = len(match.group(1))
                del headings[level - 1 :]
                headings.append(match.group(2).strip())
            else:
                current.append(line)
        flush()
        return blocks


def get_document_processor(name: str) -> DocumentProcessor:
    factories: dict[str, type[DefaultDocumentProcessor]] = {
        "default": DefaultDocumentProcessor,
    }
    try:
        return factories[name]()
    except KeyError as exc:
        raise UnsupportedDocumentError(
            f"unsupported_document_processor:{name}"
        ) from exc
