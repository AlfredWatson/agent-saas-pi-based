import asyncio
from io import BytesIO

import pytest
from app.core.config import Settings
from app.domain.rag.chunking import split_documents, token_count
from app.domain.rag.graph import (
    canonical_key,
    merge_property_maps,
    reciprocal_rank_fusion,
)
from app.integrations.rag import model_clients, processors
from app.api.v1.rag.knowledge_bases import sanitized_stored_filename
from app.db.rag.models import RagDocument
from app.db.rag.startup import _migration_head
from app.integrations.rag.file_storage import PostgresFileStorage
from app.integrations.rag.vector_store import collection_name
from app.integrations.rag.processors import (
    DefaultDocumentProcessor,
    UnsupportedDocumentError,
)
from app.workers.rag.handlers import stage_handlers
from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from openpyxl import Workbook
from pptx import Presentation


class FakeEmbeddings(Embeddings):
    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[float(len(text)), float(sum(map(ord, text)) % 101)] for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self.embed_documents([text])[0]


def test_markdown_parser_preserves_heading_path():
    documents = DefaultDocumentProcessor().parse(
        "guide.md", b"# A\nfirst\n## B\nsecond"
    )
    assert [item.page_content for item in documents] == ["first", "second"]
    assert documents[1].metadata["heading"] == "A / B"


def test_pdf_parser_preserves_page_number(monkeypatch):
    class Page:
        def extract_text(self):
            return "PDF body"

    class Reader:
        pages = [Page()]

    monkeypatch.setattr(processors, "PdfReader", lambda _: Reader())
    documents = DefaultDocumentProcessor().parse("sample.pdf", b"pdf")
    assert documents[0].page_content == "PDF body"
    assert documents[0].metadata["page"] == 1


def test_office_parsers_extract_text():
    word_buffer = BytesIO()
    word = WordDocument()
    word.add_heading("Title", 1)
    word.add_paragraph("Word body")
    word.save(word_buffer)
    assert (
        DefaultDocumentProcessor()
        .parse("sample.docx", word_buffer.getvalue())[0]
        .page_content
        == "Word body"
    )

    excel_buffer = BytesIO()
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["alpha", 2])
    workbook.save(excel_buffer)
    excel = DefaultDocumentProcessor().parse("sample.xlsx", excel_buffer.getvalue())
    assert excel[0].metadata["sheet"] == "Data"
    assert excel[0].page_content == "alpha\t2"

    ppt_buffer = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Slide title"
    presentation.save(ppt_buffer)
    powerpoint = DefaultDocumentProcessor().parse("sample.pptx", ppt_buffer.getvalue())
    assert powerpoint[0].metadata["slide"] == 1
    assert powerpoint[0].page_content == "Slide title"


@pytest.mark.parametrize("extension", ["doc", "xls", "ppt"])
def test_legacy_office_parser_extracts_single_byte_and_unicode_runs(extension):
    payload = b"binary-prefix\x00Legacy text value\x00" + "中文内容".encode("utf-16le")
    documents = DefaultDocumentProcessor().parse(f"sample.{extension}", payload)
    assert "Legacy text value" in documents[0].page_content
    assert "中文内容" in documents[0].page_content
    assert documents[0].metadata["legacy_format"] == extension


def test_unsupported_document_is_rejected():
    with pytest.raises(UnsupportedDocumentError, match="unsupported_document_type"):
        DefaultDocumentProcessor().parse("archive.zip", b"data")


def test_fixed_and_regex_chunking_enforce_token_bound():
    source = [Document(page_content=("one two three. " * 200), metadata={"page": 1})]
    fixed = split_documents(
        "fixed",
        {"max_token_size": 64, "overlap_token_size": 8, "split_by_character": "\n\n"},
        source,
    )
    regex = split_documents(
        "regex",
        {"max_token_size": 64, "re_expression": r"(?<=[.!?])\s+"},
        source,
    )
    assert len(fixed) > 1 and len(regex) > 1
    assert max(token_count(item.page_content) for item in fixed) <= 64
    assert max(token_count(item.page_content) for item in regex) <= 64


def test_semantic_chunking_requires_embeddings():
    source = [Document(page_content="First sentence. Second sentence. Third sentence.")]
    config = {"max_token_size": 64, "breakpoint_threshold": 95}
    with pytest.raises(ValueError, match="embedding_model_required"):
        split_documents("semantic", config, source)
    assert split_documents("semantic", config, source, FakeEmbeddings())


def test_graph_normalization_property_merge_and_rrf_are_deterministic():
    assert canonical_key("Company", "ＡＣＭＥ, Inc.") == "company:acme inc"
    assert merge_property_maps({"country": "CN"}, {"country": "US", "size": 2}) == {
        "country": ["CN", "US"],
        "size": 2,
    }
    assert reciprocal_rank_fusion([(["b", "a"], 1), (["a", "b"], 1)], rrf_k=60) == [
        ("a", pytest.approx(1 / 61 + 1 / 62)),
        ("b", pytest.approx(1 / 61 + 1 / 62)),
    ]


def test_rag_backend_and_concurrency_settings_fail_fast():
    with pytest.raises(ValueError, match="FILE_BASE"):
        Settings(file_base="pm")
    assert Settings(vector_base="pmcq").enabled_vector_backends == (
        "postgresql",
        "milvus",
        "chroma",
        "qdrant",
    )
    with pytest.raises(ValueError, match="VECTOR_BASE"):
        Settings(vector_base="pn")
    with pytest.raises(ValueError, match="GRAPH_BASE"):
        Settings(graph_base="pn")
    with pytest.raises(ValueError, match="default chunking concurrency"):
        Settings(rag_default_chunking_concurrency=9, rag_max_chunking_concurrency=8)
    with pytest.raises(ValueError, match="REDIS_PASSWORD"):
        Settings(
            environment="production",
            jwt_secret="a-secure-jwt-secret-value-that-is-not-a-placeholder",
            runtime_shared_secret="a-secure-runtime-secret",
            encryption_key="MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=",
            postgres_password="database-secret",
            redis_password="replace-with-a-long-redis-password",
        )


def test_external_vector_collection_name_is_deterministic_and_safe():
    from uuid import UUID

    name = collection_name(
        Settings(rag_vector_collection_prefix="pi_saas_rag"),
        UUID("12345678-1234-5678-1234-567812345678"),
    )
    assert name == "pi_saas_rag_12345678123456781234567812345678"


def test_persisted_filename_uses_user_uuid_utc_date_and_safe_stem():
    from uuid import UUID

    user_id = UUID("12345678-1234-5678-1234-567812345678")
    original, stored, extension = sanitized_stored_filename(
        user_id, "../../unsafe name?.PDF"
    )
    assert original == "unsafe name?.PDF"
    assert stored.startswith(f"{user_id}_")
    assert stored.endswith("unsafe_name.pdf")
    assert extension == ".pdf"


def test_postgres_file_storage_hides_document_bytea_access():
    source = RagDocument(storage_key="source", storage_backend="postgresql")
    target = RagDocument(storage_key="target", storage_backend="postgresql")
    storage = PostgresFileStorage(None)  # type: ignore[arg-type]

    async def scenario():
        await storage.put(source, b"content")
        assert await storage.read(source) == b"content"
        await storage.copy(source, target)
        assert await storage.read(target) == b"content"
        await storage.delete(source)
        with pytest.raises(RuntimeError, match="content_missing"):
            await storage.read(source)

    asyncio.run(scenario())


def test_redis_url_quotes_acl_credentials():
    settings = Settings(redis_username="admin-user", redis_password="secret@value")
    assert settings.redis_url.startswith("redis://admin-user:secret%40value@")
    assert settings.redis_display_url == "redis://admin-user:***@127.0.0.1:6379/0"
    assert "secret" not in settings.redis_display_url


def test_model_base_url_rejects_credentials_and_private_addresses(monkeypatch):
    monkeypatch.setattr(
        model_clients,
        "get_settings",
        lambda: Settings(rag_model_base_url_allow_private=False),
    )
    with pytest.raises(ValueError, match="invalid_model_base_url"):
        asyncio.run(model_clients.validate_base_url("https://user:secret@example.com"))
    with pytest.raises(ValueError, match="private_model_base_url_forbidden"):
        asyncio.run(model_clients.validate_base_url("https://127.0.0.1/v1"))


def test_stage_handler_registry_is_explicit_and_extensible():
    class Worker:
        async def process_parsing(self, db, job):
            return None

        async def process_chunking(self, db, job):
            return None

        async def process_vectorization(self, db, job):
            return None

        async def process_graph(self, db, job):
            return None

    assert set(stage_handlers(Worker())) == {
        "parsing",
        "chunking",
        "vectorization",
        "graph_extraction",
    }


def test_rag_migration_head_resolves_from_the_layered_module_location():
    assert _migration_head() == "0010_workspace_agent_workbench"
